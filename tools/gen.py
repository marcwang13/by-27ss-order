import sys, json, hashlib, io, os, re
sys.path.insert(0, os.path.dirname(__file__))
from parse import load
from collections import OrderedDict, defaultdict
FX = 0.0245
def twd(k): return int(round(k*FX/10.0))*10 if k else None
CAT = {'Crew-neck T-Shirts':'圓領T恤','Polo Shirts':'POLO衫','Zip T-Shirts':'拉鍊衫','Other T-Shirts':'其他T恤','Shirts':'襯衫','Hoodies':'連帽衫',
 'Jackets':'外套','GORE-TEX Jackets':'GORE-TEX 外套','Padded Jackets':'鋪棉外套','Set-up Jackets':'套裝外套','Vests':'背心','Parkas':'連帽外套',
 'Pants':'長褲','Shorts':'短褲','Set-up Pants':'套裝長褲','Accessory Bags':'配件包','Backpacks':'背包','Gloves':'手套','Headwear':'帽款',
 'Lighting Accessories':'照明配件','Other Equipment':'其他裝備','Socks':'襪子','Sunglasses':'太陽眼鏡','Trekking Poles':'登山杖',
 '후드티셔츠':'連帽衫','셋업':'套裝','세트':'套組'}
SHEET = {}
for k in ['圓領T恤','POLO衫','拉鍊衫','其他T恤','襯衫','連帽衫']: SHEET[k]='上衣'
for k in ['外套','GORE-TEX 外套','鋪棉外套','套裝外套','背心','連帽外套']: SHEET[k]='外套'
for k in ['長褲','短褲','套裝長褲']: SHEET[k]='褲裝'
for k in ['套裝','套組']: SHEET[k]='套裝'
for k in ['配件包','背包','手套','帽款','照明配件','其他裝備','襪子','太陽眼鏡','登山杖']: SHEET[k]='配件'
LINE = {'알파인A':'ALPINE A','터프쉴드':'TOUGH SHIELD','트레일러닝':'TRAIL RUNNING','클라이밍':'CLIMBING'}
PTYPE = {'일반':'一般','라운드넥':'圓領','경량':'輕量','중헤비':'中厚','집업':'拉鍊','후디':'連帽','방풍':'防風','플리스':'刷毛','충전재':'填充保暖','카라':'有領','레깅스팬츠':'內搭褲','2L':'2L 防水'}
LETTER = ['0XS','00S','00M','00L','0XL','XXL','3XL']
SZLABEL = {'0XS':'XS','00S':'S','00M':'M','00L':'L','0XL':'XL','XXL':'2XL','3XL':'3XL','999':'F','ASSORTED':'數量'}
def szkey(z):
    if z in LETTER: return (0, LETTER.index(z))
    if z=='999': return (3,0)
    try: return (1 if int(z)>=85 or int(z)<30 else 2, int(z))
    except: return (4,0)
def clean(s):
    if s is None: return ''
    s=str(s).strip()
    return '' if re.search(r'pending|^TBD$|^-$', s, re.I) else s
def gender_ss(name, style):
    n=' '+(name or '').upper()+' '
    if ' WOMEN' in n or n.startswith(' W '): return '女'
    if ' MEN' in n: return '男'
    if ' UNI' in n: return '中性'
    m=re.search(r'(\d)\d\d$', style or '')
    if m: d=int(m.group(1)); return '男' if d<5 else ('女' if d<9 else '中性')
    return '中性'

NAMES={}
def build(season, xls):
    rows=load(xls)['Order Sheet']; h=rows[0]; R=[dict(zip(h,r)) for r in rows[1:]]
    items=OrderedDict(); runs=defaultdict(set)
    for r in R:
        style=r['Product Code']; cc=clean(r['Color Code']); size=r['Size'] or 'ASSORTED'
        typ=r['Workbook Variety'] if season=='27FW' else r['Product Type']
        if season=='27FW' and typ in ('-',None): typ=r['Product Type']
        cat=CAT.get(typ, PTYPE.get(typ, typ or '其他'))
        if season=='27FW' and cat in PTYPE.values() and typ not in CAT: cat=cat
        sheet=SHEET.get(cat,'上衣')
        if season=='27FW':
            key=style; sizes=['ASSORTED']
        else:
            key=f'{style}-{cc}'
        krw=int(float(r['KR SRP'])) if r['KR SRP'] else None
        if key not in items:
            en=clean(r['Product Name (English)'])
            if season=='27SS':
                inb=r['Yongin Inbound'] or ''
                mon=(str(int(inb[5:7]))+'月') if inb else '未定'
                line=r['Product Line'] or ''
                gen=gender_ss(en, style)
                color=clean(r['Color Name']); 
                note=[]
                if r['Category']=='Special-price Apparel': note.append('特價商品')
                if r['Target Nation']: note.append('原廠標示：'+r['Target Nation'])
                en=en or NAMES.get(style,'')
                gen=gender_ss(en, style)
                it=dict(id=key, sheet=sheet, no='', mon=mon, monTag=f'{mon}韓國入倉', theme=line, drop=r['Pricing Class'], cat=cat, gen=gen,
                    style=style, zh='', en=en or '英文品名待原廠提供', color=color or '顏色待確認', cc=cc, sku=style+cc, icon=clean(r['Special Material']),
                    fab=clean(r['Composition']), series=clean(r['Country of Origin']) and ('產地 '+clean(r['Country of Origin'])), ov=clean(r['Product Description']), dt='', tech='',
                    note='；'.join(note), launch=inb.replace('-','/'), eta=inb.replace('-','/'), krw=krw, price=twd(krw), sizes=[], doll='', run=style, part='top')
            else:
                line=LINE.get(r['Product Line'], r['Product Line'] or '')
                st=style.upper(); g={'M':'男','W':'女','U':'中性'}.get(st[3:4],'中性')
                szr=clean(r['Size']); szr='' if szr=='ASSORTED' else szr
                color=clean(r['Color Name']); color='' if color=='ASSORTED' else color
                ptype='' if r['Product Type'] in ('-',None) else PTYPE.get(r['Product Type'], r['Product Type'])
                note=[]
                if szr: note.append('尺寸段：'+szr.replace('\n','／'))
                if r['Target Nation']: note.append('原廠標示：'+r['Target Nation'])
                it=dict(id=key, sheet=sheet, no='', mon=line or '其他', monTag=line or '', theme=line, drop=r['Pricing Class'], cat=cat, gen=g,
                    style=style, zh='', en=en or '品名待原廠提供', color=color or '多色（色別待原廠提供）', cc='', sku=style, icon=clean(r['Special Material']) or ptype,
                    fab=clean(r['Composition']), series=clean(r['Country of Origin']) and ('產地 '+clean(r['Country of Origin'])), ov=clean(r['Product Description']), dt='', tech='',
                    note='；'.join(note), launch='', eta='', krw=krw, price=twd(krw), sizes=[], doll='', run=style, part='top')
            items[key]=it
        it=items[key]
        if size not in it['sizes']: it['sizes'].append(size)
        runs[style].add(size)
    for it in items.values(): it['sizes'].sort(key=szkey)
    order=['上衣','外套','褲裝','套裝','配件']
    if season=='27SS': mo=['12月','1月','2月','3月','4月','未定']
    else:
        mo=['TOUGH SHIELD','ALPINE A','ALPINE B','TRAIL RUNNING','CLIMBING','SKI']
    lst=list(items.values()); idx={id(x):n for n,x in enumerate(lst)}
    lst.sort(key=lambda x:(order.index(x['sheet']), mo.index(x['mon']) if x['mon'] in mo else 99, idx[id(x)]))
    items=OrderedDict((x['id'],x) for x in lst)
    RUNS={k:sorted(v,key=szkey) for k,v in runs.items()}
    return list(items.values()), RUNS

def write(season, items, RUNS, outdir, imgmap):
    os.makedirs(outdir+'/big', exist_ok=True)
    from PIL import Image
    IMG={}
    for it in items:
        src=imgmap.get(it['style'])
        if not src: continue
        hsh=hashlib.md5(src).hexdigest()[:12]
        if hsh not in IMG:
            im=Image.open(io.BytesIO(src)).convert('RGB')
            b=im.copy(); b.thumbnail((1200,1200)); b.save(f'{outdir}/big/{hsh}.jpg',quality=85)
            t=im.copy(); t.thumbnail((360,360)); bio=io.BytesIO(); t.save(bio,'JPEG',quality=72)
            import base64; IMG[hsh]='data:image/jpeg;base64,'+base64.b64encode(bio.getvalue()).decode()
        it['doll']=hsh
    order=['上衣','外套','褲裝','套裝','配件']
    sheets=[s for s in order if any(i['sheet']==s for i in items)]
    if season=='27SS':
        mons=['12月','1月','2月','3月','4月','未定']; mons=[m for m in mons if any(i['mon']==m for i in items)]
    else:
        mons=[]; [mons.append(i['mon']) for i in items if i['mon'] not in mons]
    sorder=sorted({z for v in RUNS.values() for z in v}, key=szkey)
    cfg=dict(season=season, fx=FX, sheets=sheets, months=mons, szOrder=sorder, szLabel=SZLABEL)
    with open(outdir+'/data.js','w') as f:
        f.write('window.SDR_ITEMS='+json.dumps(items,ensure_ascii=False)+';\n')
        f.write('window.BY_RUNS='+json.dumps(RUNS,ensure_ascii=False)+';\n')
        f.write('window.BY_CFG='+json.dumps(cfg,ensure_ascii=False)+';\n')
    with open(outdir+'/imgs.js','w') as f: f.write('window.SDR_IMG='+json.dumps(IMG)+';\n')
    return sheets, mons, len(IMG)

if __name__=='__main__':
    season, xls, outdir = sys.argv[1:4]
    imgmap={}
    if len(sys.argv)>4:
        import openpyxl
        wb=openpyxl.load_workbook(sys.argv[4]); ws=wb.worksheets[0]
        for im in ws._images:
            r=im.anchor._from.row+1; st=ws.cell(r,7).value
            if st and st not in imgmap: imgmap[st]=im._data()
        for r in range(9, ws.max_row+1):
            st=ws.cell(r,7).value; nm=ws.cell(r,8).value
            if st and nm and 'pending' not in str(nm).lower(): NAMES.setdefault(st, str(nm).strip())
    items,RUNS=build(season,xls)
    print(season, 'items',len(items), 'styles',len(RUNS), write(season,items,RUNS,outdir,imgmap))
    from collections import Counter
    print(Counter(i['sheet'] for i in items), Counter(i['gen'] for i in items), 'noimg', sum(1 for i in items if not i['doll']))
